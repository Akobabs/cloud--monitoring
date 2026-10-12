"""Write the retrained research continuation from audited experiment artifacts.

Run using the managed document Python runtime. Preserves prior Word editions.
"""
import argparse
import json
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_research_document import Manuscript, ROOT, SHORT, fmt


class RetrainedManuscript(Manuscript):
    def table(self, caption, headers, rows, widths=None):
        super().table(caption, headers, rows, widths)
        self.doc.paragraphs[-2].paragraph_format.keep_with_next = True
        table = self.doc.tables[-1]
        if len(table.rows) <= 9:
            for index, row in enumerate(table.rows):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", type=Path, required=True)
    args = parser.parse_args()
    directory = args.run.resolve()
    read = lambda name: json.loads((directory / name).read_text(encoding="utf-8"))
    run, config = read("run.json"), read("config.json")
    audit, comparison_audit = read("evidence_audit.json"), read("comparison_audit.json")
    if run["status"] != "complete" or not run.get("experiment_revision"):
        raise ValueError("A complete retrained experiment is required")
    if audit["status"] != "passed" or comparison_audit["status"] != "passed":
        raise ValueError("Both evidence audits must pass")
    summary = pd.read_csv(directory / "summary.csv")
    metrics = pd.read_csv(directory / "metrics.csv")
    comparison = pd.read_csv(directory / "improvement_comparison.csv")
    datasets = {p.parent.name: json.loads(p.read_text()) for p in directory.glob("datasets/*/dataset.json")}
    manifest = read("data_manifest.json")
    screenshots = directory / "screenshots"
    ui = json.loads((screenshots / "ui_review.json").read_text())
    if ui["browser_errors"]:
        raise ValueError("Resolve browser errors before building the document")
    google = summary[summary.family == "Google failure proxy"].sort_values("f1", ascending=False)
    smd = summary[summary.family == "SMD adapted benchmark"].sort_values("f1", ascending=False)
    best_g, best_s = google.iloc[0], smd.iloc[0]
    wins = int((comparison.f1_change > 1e-12).sum())
    losses = int((comparison.f1_change < -1e-12).sum())
    equal = len(comparison) - wins - losses
    candidate_fits = sum(len(json.loads(p.read_text())) for p in directory.glob("models/*/*/*/tuning.json"))
    g = datasets["google_cluster"]
    positive = g["eligible_partitions"]["test"]
    prevalence = positive["positive"] / positive["rows"]
    rf = comparison[(comparison.family == "Google failure proxy") & (comparison.model_name == "Random Forest")].iloc[0]
    sif = comparison[(comparison.family == "SMD adapted benchmark") & (comparison.model_name == "Isolation Forest")].iloc[0]
    m = RetrainedManuscript()
    title = "Comparative Analysis of Anomaly Detection for Cloud Infrastructure Using Selected Machine Learning Algorithms"
    m.doc.add_paragraph(title, "Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.markdown.append("# " + title + "\n")
    m.paragraph("Implementation plan and research continuation")
    m.paragraph("Chapters Four and Five with Appendices A to C")
    m.paragraph("September 2026")
    m.paragraph(f"This edition reports the retrained experiment {run['run_id']}, completed on {run['completed_utc'][:10]}, and compares it with the preserved validation-calibrated baseline {comparison_audit['baseline_run']}. All {len(metrics)} model and seed evaluations were fitted again. The original Chapters One to Three are retained separately and unchanged.")
    m.paragraph("The four-algorithm comparison uses three selected Server Machine Dataset series. Google Cluster records support a supplementary comparison of three tabular algorithms; their short task histories do not support the specified LSTM. Findings remain exploratory because the evaluation period had been reviewed before this revision.")

    m.heading("Implementation Plan", page=True)
    m.paragraph("The revision addresses two implementation choices: resource representation and consistency between candidate selection and the final decision threshold. The experiment plan and candidate grid were fixed before the revised run. Data acquisition, partitioning, common fitting endpoints and the seed schedule were retained so differences could be interpreted against the existing baseline.")
    m.table("Plan 1 Work packages and evidence", ["Work package", "Implementation", "Evidence"], [
        ["Feature engineering", "CPU and memory usage-to-request ratios", "13 Google predictors"],
        ["Candidate selection", "Calibrate each candidate before comparing validation F1", "tuning.json and selection.json"],
        ["Fresh experiment", "Three seeds and four datasets", f"{len(metrics)} evaluations; {candidate_fits} candidate fits"],
        ["Verification", "Recompute decisions and match baseline records", "Two passed evidence audits"],
        ["Reporting", "Retain gains, regressions and fresh resource costs", "Comparison CSV, figures and dashboard"],
    ], [1.15, 2.35, 2.268])
    m.paragraph("The baseline artifacts remain unchanged. The improved profile is reproduced with run_project.bat improved, and the local dashboard is opened with run_frontend.bat. The configuration, raw-file hashes, dependency versions and source fingerprint are retained in the new run directory. No separate unit-test suite was introduced; verification uses execution and artifact audits.")

    m.heading("CHAPTER FOUR", page=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("IMPLEMENTATION AND RESULTS").alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("4.1 Introduction", 2)
    m.paragraph("The framework compares Isolation Forest, One-Class SVM, Random Forest and an LSTM Autoencoder using chronological partitions and common evaluation endpoints. F1 is the primary detection metric; precision, recall, false-positive rate, ROC-AUC and average precision provide context. Computational measurements describe offline CPU execution, not production serving performance.")
    m.heading("4.2 Implementation Environment", 2)
    m.table("Table 4.1 Recorded execution environment", ["Component", "Value"], [
        ["Operating system", run["hardware"]["platform"]],
        ["Python", run["hardware"]["python"].split()[0]],
        ["Memory", f"{run['hardware']['ram_gib']:.2f} GiB"],
        ["Scikit-learn", run["versions"]["scikit-learn"]],
        ["TensorFlow and Keras", run["versions"]["tensorflow"] + " and " + run["versions"]["keras"]],
        ["Execution", f"CPU; {config['threads']} compute threads; sequential model processes"],
        ["Random seeds", ", ".join(map(str, config["seeds"]))],
    ], [1.65, 4.118])
    m.paragraph("Deterministic TensorFlow operations and random seeds were requested. Such controls do not guarantee bitwise equivalence across library or hardware changes. The environment snapshot is the authoritative dependency record for this run.")
    m.heading("4.3 Methodological Amendments", 2)
    m.paragraph("The earlier baseline selected model settings using F1 at a fixed percentile threshold and then recalibrated only the selected model. The revised implementation calibrates every candidate on labelled validation data before selecting the candidate with the highest validation F1. This aligns the selection criterion with the final classification rule. It uses the same candidate grid and fitting budget as the reference experiment.")
    m.paragraph("The second change adds two row-local Google resource ratios. All models were refitted, so weights, continuous scores and resource measurements may change. This is a modelling experiment, unlike the earlier threshold-only follow-up. Both modifications were applied together on Google; their separate causal contributions cannot be identified without a further ablation. SMD retains its original features and evaluates the changed selection procedure.")
    m.heading("4.4 Dataset Acquisition and Audit", 2)
    m.paragraph(f"The Google source is Kaggle version 1 of the Google 2019 Cluster sample. Its {g['source_rows']:,} records included {g['duplicate_telemetry_rows_removed']:,} repeated telemetry rows. Deduplication preceded chronological partitioning, and {g['cross_partition_task_rows_purged']:,} rows were purged because their task crossed a partition boundary. Failure labels agree with event=FAIL. A zero label includes other event types and is not a verified healthy-operation label.")
    m.paragraph("SMD supplies ordered, labelled metric records. This study re-partitions the published labelled test sequence into 60/20/20 chronological segments to permit supervised Random Forest training. The official normal-only training files are downloaded and hashed but not fitted. This adaptation is not directly comparable with the official SMD protocol. Only machine-1-1, machine-2-1 and machine-3-1 are included.")
    rows = []
    for info in datasets.values():
        for part, counts in info["eligible_partitions"].items():
            rows.append([info["display_name"].replace("Google Cluster 2019 sample", "Google").replace("SMD machine-", "SMD "), part, f"{counts['rows']:,}", f"{counts['positive']:,}"])
    m.table("Table 4.2 Eligible endpoints after partitioning and sequence warm-up", ["Dataset", "Partition", "Rows", "Positives"], rows, [1.45, 1.45, 1.45, 1.418])
    m.paragraph(f"Google tasks contain at most {g['max_observations_per_task_after_deduplication']} distinct observations after deduplication. Joining unrelated tasks into a 30-step sequence would create artificial histories, so the Google LSTM remains not applicable. Google outcome timing is not independently established; the task is retrospective failure-proxy discrimination, not demonstrated prospective prediction.")
    m.heading("4.5 Preprocessing and Feature Engineering", 2)
    m.paragraph("The split precedes fitted transformations. Missing values are filled using training medians; entirely unobserved training columns are excluded. Standard scaling is fitted on training data and applied without refitting to validation and test partitions. Infinite engineered values are treated as missing. Input checks reject non-binary SMD labels before integer conversion, inconsistent Google labels, missing telemetry keys and invalid usage times.")
    m.paragraph("Google retains its 11 resource predictors and adds CPU utilisation = average_cpu / requested_cpu and memory utilisation = average_memory / requested_memory. A ratio is missing when the request is non-positive and is then imputed from training values. No arbitrary epsilon is inserted into the denominator. These features describe resource use relative to allocation and contain no outcome flags, identifiers, timestamps or observations from another row.")
    m.paragraph("SMD retains 38 anonymous metric channels, their trailing five-observation means and first differences, giving 114 inputs. Rolling calculations restart at each partition boundary. All four algorithms are evaluated at the same endpoints after excluding the first 29 records of each partition; LSTM inputs are genuine 30-step windows. No physical channel names or sampling interval are invented.")
    m.heading("4.6 Model Training and Selection", 2)
    m.table("Table 4.3 Fixed candidate settings", ["Algorithm", "Candidates"], [
        ["Isolation Forest", "100 or 200 trees; max_samples 256; normal windows only"],
        ["One-Class SVM", "RBF kernel; nu 0.01 or 0.05; gamma scale; normal windows only"],
        ["Random Forest", "100 trees/depth 10 or 200 trees/depth 20; balanced class weights; leaf minimum 2"],
        ["LSTM Autoencoder", "Two encoder and two decoder layers; latent width 8 or 16; Adam; MSE; up to 15 epochs; batch 64; patience 3"],
    ], [1.45, 4.318])
    m.paragraph("Each seed samples up to 6,000 shared training endpoints. Random Forest uses both labels; the other algorithms fit only endpoints whose entire input window is normal. Thus the study compares learning regimes as well as model families. LSTM early stopping uses up to 2,000 entirely normal validation windows; its score is reconstruction MSE at the last step.")
    m.heading("4.7 Evaluation Procedure", 2)
    m.paragraph("Each candidate is scored on validation records. When both labels are present, its threshold maximises validation F1, with the higher threshold preferred on exact ties. Predictions use score > threshold consistently. Candidate ties retain the first predefined setting. Where validation lacks a class, the percentile of entirely normal validation scores is retained; machine-3-1 has no validation positives and uses the first predefined candidate. Selected settings, thresholds and validation scores are saved before test scoring.")
    m.paragraph("F1, precision and recall are calculated on raw point predictions with no point adjustment. ROC-AUC and average precision use continuous scores. SMD family results first macro-average machines within each seed, then report means and sample standard deviations across seeds. These deviations do not establish statistical significance or uncertainty across unseen environments. The reused test period makes this an exploratory follow-up even though no test labels enter candidate or threshold selection.")
    m.heading("4.8 Detection Results and Baseline Comparison", 2)
    m.table("Table 4.4 F1 before and after retraining", ["Dataset", "Algorithm", "Before", "After", "Change"], [
        ["SMD" if r.family.startswith("SMD") else "Google", SHORT[r.model_name], fmt(r.f1_baseline), fmt(r.f1_revised), f"{r.f1_change:+.4f}"] for r in comparison.itertuples()
    ], [.6, 1.6, 1.15, 1.15, 1.268])
    m.paragraph(f"Across the seven dataset-family and algorithm combinations, F1 increased in {wins}, decreased in {losses}, and was unchanged in {equal}. The audit identified {comparison_audit['changed_candidate_selections']} changed candidate selections among {len(metrics)} evaluations. No model or seed was excluded because of a regression.")
    m.paragraph(f"Google Random Forest F1 changed from {rf.f1_baseline:.4f} to {rf.f1_revised:.4f}; its false-positive rate changed from {rf.false_positive_rate_baseline:.4f} to {rf.false_positive_rate_revised:.4f}. SMD Isolation Forest F1 changed from {sif.f1_baseline:.4f} to {sif.f1_revised:.4f}, with false-positive rate changing from {sif.false_positive_rate_baseline:.4f} to {sif.false_positive_rate_revised:.4f}. These are measured improvements in the selected operating points, not evidence of broad generalisation.")
    m.table("Table 4.5 False-positive rates before and after retraining", ["Dataset", "Algorithm", "Before", "After"], [
        ["SMD" if r.family.startswith("SMD") else "Google", SHORT[r.model_name], fmt(r.false_positive_rate_baseline), fmt(r.false_positive_rate_revised)] for r in comparison.itertuples()
    ], [.7, 2.05, 1.5, 1.518])
    for number, label, frame in [(6, "SMD adapted benchmark", smd), (7, "Google failure proxy", google)]:
        result_rows = [[SHORT[r.model_name], fmt(r.precision), fmt(r.recall), fmt(r.f1), fmt(r.roc_auc), fmt(r.false_positive_rate)] for r in frame.itertuples()]
        if label.startswith("Google"):
            result_rows.append(["LSTM AE", "N/A", "N/A", "N/A", "N/A", "N/A"])
        m.table(f"Table 4.{number} Retrained detection results for {label}", ["Algorithm", "Precision", "Recall", "F1", "ROC-AUC", "FPR"], result_rows, [1.5, .85, .8, .8, .95, .868])
    m.paragraph(f"The highest revised mean F1 was {best_s.f1:.4f} for {best_s.model_name} on SMD and {best_g.f1:.4f} for {best_g.model_name} on Google. A trivial all-positive Google rule obtains F1 {2 * prevalence / (1 + prevalence):.4f} with recall and false-positive rate both equal to 1. The Google findings must be assessed against this reference and their still substantial false-alarm burden.")
    m.table("Table 4.8 Seed variation and score ranking", ["Dataset", "Algorithm", "F1 SD", "Average precision"], [
        ["SMD" if r.family.startswith("SMD") else "Google", SHORT[r.model_name], fmt(r.f1_sd), fmt(r.average_precision)] for r in summary.itertuples()
    ], [.7, 2.05, 1.2, 1.818])
    per_machine = metrics[metrics.dataset.str.startswith("smd_")].groupby(["dataset", "model_name"]).f1.mean().unstack()
    m.table("Table 4.9 Mean F1 by SMD machine", ["Machine", "Isolation Forest", "One-Class SVM", "Random Forest", "LSTM AE"], [
        [name.removeprefix("smd_machine-"), *[fmt(row[model]) for model in SHORT]] for name, row in per_machine.iterrows()
    ], [.65, 1.35, 1.3, 1.3, 1.168])
    m.paragraph("Machine-1-1 has only seven positive test endpoints, while machine-3-1 has no positive validation endpoints. These distributions explain why a threshold selected earlier may transfer poorly. Aggregate results must be read alongside the machine-level table.")
    m.figure(directory / "figures/smd_comparison.png", "Figure 4.1 Retrained SMD F1 and false-positive rates with seed standard deviation")
    m.figure(directory / "figures/google_comparison.png", "Figure 4.2 Retrained Google F1 and false-positive rates with seed standard deviation")
    m.heading("4.9 Computational Results", 2)
    m.table("Table 4.10 Newly measured computational costs", ["Dataset", "Algorithm", "Fit seconds", "ms per point", "Peak RSS MB"], [
        ["SMD" if r.family.startswith("SMD") else "Google", SHORT[r.model_name], fmt(r.train_seconds, 2), fmt(r.latency_ms_per_observation), fmt(r.peak_rss_mb, 1)] for r in summary.itertuples()
    ], [.6, 1.6, 1.15, 1.15, 1.268])
    m.paragraph("All resource values were measured during this new run. Fit time is the selected candidate's fit call; it excludes preprocessing and model construction. Tuning time is separately recorded. Inference is warmed and timed three times, with median batch time divided by the observation count. RSS is sampled every 10 ms and includes library and data overhead; brief peaks may be missed. These measurements should not be interpreted as single-request serving latency or exact per-model allocations.")
    m.heading("4.10 Interface and Verification", 2)
    m.paragraph(f"The evidence audit verified all {audit['evaluations']} saved evaluations across {audit['datasets']} datasets, including confusion counts, F1, reproducible validation decisions, aligned endpoints, finite features, partition separation and the source fingerprint. A second audit confirmed identical test row identifiers, labels, original prepared feature values and sampled training endpoints against the baseline. Only the two intended Google columns were added.")
    m.paragraph("The dashboard reads saved artifacts, displays before-and-after F1 and false-positive rates, and retains access to older runs. All five views were exercised and a results CSV download was checked without browser errors. Screenshots in this edition were captured from the completed retrained experiment. The interface does not retrain models or issue production alerts.")
    m.figure(screenshots / "01_overview.png", "Figure 4.3 Dashboard showing the retrained run and baseline comparison")
    m.heading("4.11 Discussion and Limitations", 2)
    m.paragraph("Aligning candidate selection with the final threshold rule corrects an optimisation mismatch, but it does not guarantee higher test F1. The One-Class SVM regressions are retained as evidence of this limitation. Resource ratios provide a representation of allocation pressure, but the combined Google experiment does not isolate their effect from candidate selection. Validation is reused for tuning and calibration, and for LSTM early stopping; this may overfit validation decisions.")
    m.paragraph("The Google false-positive rates remain too high to justify an operational deployment claim. The best revised Google model must also be compared with the strongest earlier model, rather than only its own baseline. Small numerical gains across three seeds are not proof of statistical significance. No additional test-driven tuning round was performed after this comparison.")
    m.paragraph("The experiment uses a reviewed evaluation period, three selected SMD machines, an adapted supervised split and a bounded candidate grid. It does not establish prospective failure prediction, robustness to drift, incident-level detection delay or production cost. Fresh environments and evaluation periods are required before generalisation can be assessed.")
    m.heading("4.12 Chapter Summary", 2)
    m.paragraph(f"The revision produced {len(metrics)} fresh model evaluations and an auditable before-and-after comparison. It changed both learned scores and candidate choices where supported by validation. The outcomes include {wins} F1 gains, {losses} regressions and {equal} unchanged family-level result. Reporting these differences together with false alarms and resource costs provides the basis for the conclusions.")

    m.heading("CHAPTER FIVE", page=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("SUMMARY CONCLUSION AND RECOMMENDATION").alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("5.1 Summary of Study and Findings", 2)
    m.paragraph("The study implemented and compared four cloud telemetry anomaly-detection algorithms. Chronological preprocessing, shared endpoints, seed repetition and separate dataset interpretations support a controlled offline comparison. Google remains a supplementary tabular failure-proxy task because its task histories cannot support the requested sequence model.")
    m.paragraph(f"The retrained experiment added resource-utilisation features and calibrated every candidate before selection. Google Random Forest achieved F1 {rf.f1_revised:.4f} against {rf.f1_baseline:.4f} previously, with a lower false-positive rate. SMD Isolation Forest achieved F1 {sif.f1_revised:.4f} against {sif.f1_baseline:.4f}. Across all seven family-level comparisons, {wins} improved, {losses} regressed and {equal} remained unchanged. The original experiment was preserved and resource measurements were repeated.")
    m.heading("5.2 Conclusion", 2)
    m.paragraph(f"Within the revised experiment, {best_s.model_name} had the highest mean SMD F1 and {best_g.model_name} had the highest mean Google F1. The changes produced selective, modest gains rather than a universal improvement. A validation-selected operating point can reduce false alarms or improve F1, yet still transfer poorly to a later period. The study therefore supports careful validation and explicit reporting of regressions, not a claim that one algorithm is universally superior.")
    m.paragraph("The framework meets the implementation and offline comparison objectives. High false-alarm rates, uncertain Google label timing and reuse of an already examined test period limit practical conclusions. These findings require independent confirmation before deployment decisions.")
    m.heading("5.3 Recommendation for Future Research", 2)
    m.paragraph("Future work should separate the effects of resource features and candidate calibration through an ablation, reserve a distinct calibration segment, and evaluate the resulting choices on a new untouched time period. Rolling temporal validation would provide a stronger view of threshold stability and changing anomaly prevalence. A predeclared false-alarm budget should accompany F1 when operational alert volume matters.")
    m.paragraph("The dataset coverage should expand to all SMD machines and independent infrastructure traces. Google usage records should be linked to verified future event outcomes with explicit prediction horizons and sufficient within-task history. A separate official-split SMD experiment should distinguish normal-only learning from supervised learning. Production studies should then measure drift, event-level delay, end-to-end latency and retraining cost.")

    m.heading("APPENDIX A", page=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("CODEBASE AND REPRODUCTION").alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.paragraph("The current implementation is supplied as runnable source. run_project.bat improved executes configs/improved.json, preserving the original research and quick profiles. scripts/audit_run.py validates saved decisions; scripts/compare_experiments.py verifies baseline comparability and writes improvement_comparison.csv. scripts/build_improved_document.py creates this continuation from those audited results. Historical runs cannot be resumed with changed source because their fingerprints are intentionally enforced.")
    m.paragraph("The shared calibration routine below implements the labelled-validation threshold search. Google ratios are implemented in src/cloud_anomaly/data.py; training and selection are implemented in worker.py. The printed listings are for inspection; runnable files remain authoritative.")
    m.heading("A.1 Validation Threshold Selection", 2)
    m.code("src/cloud_anomaly/calibration.py")
    m.heading("A.2 Model Implementations", 2)
    m.code("src/cloud_anomaly/models.py")
    m.heading("A.3 Experiment Worker", 2)
    m.code("src/cloud_anomaly/worker.py")
    m.heading("APPENDIX B", page=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("INTERFACE EVIDENCE").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for file, caption in [("02_model_comparison.png", "Figure B.1 Fresh detection and resource results"), ("03_prediction_explorer.png", "Figure B.2 Held-out score and error inspection"), ("04_dataset_details.png", "Figure B.3 Dataset interpretation and partitions")]:
        m.figure(screenshots / file, caption)
    m.heading("APPENDIX C", page=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.heading("DATA PROVENANCE AND INTEGRITY").alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.paragraph(f"The Google sample is Kaggle version {manifest['google_version']}; the SMD source revision is {manifest['smd_revision']}. The Google listing identifies CC BY-NC-SA 4.0, while the original Google trace is CC BY 4.0. The SMD repository includes MIT licensing; no separate dataset-specific licence was identified. Source URLs, raw-file sizes and SHA-256 values are stored in data_manifest.json.")
    m.table("Table C.1 Google predictor definitions", ["Features", "Definition"], [
        ["average_cpu, average_memory", "Mean usage in normalised trace units"],
        ["maximum_cpu, maximum_memory", "Maximum usage in normalised trace units"],
        ["sample_cpu", "Supplied random-sample CPU usage"],
        ["requested_cpu, requested_memory", "Requested resources in normalised units"],
        ["assigned_memory, page_cache_memory", "Assigned limit and page-cache usage"],
        ["cycles_per_instruction, memory_accesses_per_instruction", "Supplied instruction-level ratios"],
        ["cpu_utilisation, memory_utilisation", "Mean usage divided by positive requested resources"],
    ], [2.6, 3.168])
    m.paragraph("SMD features are metric_01 through metric_38 and their causal mean and difference terms. Outcome flags, task identifiers and timestamps do not enter model features. Complete missingness counts, retained feature names, partition counts and exclusions are stored in each dataset.json alongside prepared.npz and the fitted preprocessor.")
    m.paragraph(f"Recorded source SHA-256: {run['source_sha256']}")
    m.paragraph(f"Recorded configuration SHA-256: {run['config_sha256']}")
    m.heading("References for the Implementation", page=True)
    for reference in [
        "Google. ClusterData 2019 traces. https://github.com/google/cluster-data/blob/master/ClusterData2019.md",
        "Mwiti, D. Google 2019 Cluster sample, version 1. https://www.kaggle.com/datasets/derrickmwiti/google-2019-cluster-sample",
        "NetManAIOps. OmniAnomaly and Server Machine Dataset. https://github.com/NetManAIOps/OmniAnomaly",
        "Scikit-learn developers. Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls.html",
        "Scikit-learn developers. Tuning the decision threshold for class prediction. https://scikit-learn.org/stable/modules/classification_threshold.html",
        "TensorFlow developers. Enable operation determinism. https://www.tensorflow.org/api_docs/python/tf/config/experimental/enable_op_determinism",
    ]:
        m.paragraph(reference)
    m.doc.core_properties.title = title
    m.doc.core_properties.subject = "Retrained implementation and measured comparison"
    m.doc.core_properties.author = ""
    output = ROOT / "docs/Research_Implementation_and_Results_Retrained.docx"
    m.doc.save(output)
    manuscript = ROOT / "docs/research_manuscript.md"
    previous = ROOT / "docs/research_manuscript_before_retraining.md"
    if manuscript.exists() and not previous.exists():
        previous.write_bytes(manuscript.read_bytes())
    manuscript.write_text("\n".join(m.markdown), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
